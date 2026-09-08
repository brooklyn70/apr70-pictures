#!/usr/bin/env python3
"""Pass 6: build the split Word document and its Markdown twin."""
import json, pathlib, re, shutil, sys, datetime
sys.path.insert(0, str(pathlib.Path(__file__).parent))
import lib

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = lib.OUT
inv = json.loads((OUT / "inventory.json").read_text())
after = json.loads((OUT / "after.json").read_text())
voice = (OUT / "voice-profile-marco.md").read_text()
flags = json.loads((OUT / "residue-flags.json").read_text()) if (OUT / "residue-flags.json").exists() else {}
repairs = json.loads((OUT / "repairs.json").read_text()) if (OUT / "repairs.json").exists() else {}
ledger = lib.ledger_load()

DEST = pathlib.Path("/Volumes/SharedData/10-03-website-assets/copy-pass-2026-09-07/APR70-site-copy-writers-room-2026-09-07.docx")
SCRATCH = pathlib.Path("/private/tmp/claude-501/-Users-marco-websites-apr70-website/2251e5c8-4793-420f-af9f-38fc6b2b3078/scratchpad/APR70-site-copy-writers-room-2026-09-07.docx")
VAULT = pathlib.Path("/Users/marco/vault/10 Work/11 APR70 Pictures/11.03 Company Ops/00 Website - Copy Pass 2026-09-07.md")

CHAIRS = [("welles", "Orson Welles", "structure, escalation, scale, the unforgettable image"),
          ("sturges", "Preston Sturges", "comic architecture, sequence mechanics, dialogue velocity"),
          ("riskin", "Robert Riskin", "the moral ledger, the working-class voice"),
          ("capra", "Frank Capra", "the audience verdict, the full-house read")]

LESSONS = [
 "The founder's voice is not a style, it is an order of operations. He puts the verdict in the first sentence and the reasoning after it. Almost every improvement in this pass came from moving the claim forward and deleting the run-up to it (\"We operate on a simple conviction\" became \"One conviction runs this company\").",
 "A voice profile built from real sent mail beats a voice brief written from memory. Thirty-five of his own emails produced testable rules (no warm-up line, source and page number on every claim, one flat sentence of consequence before the signature) that a description of his voice would never have produced.",
 "The corpus that reads worst is the one that teaches most. His dictated agent prompts are lowercase, unpunctuated and impatient, and none of that register belongs on a public page. What they carry is temperament: he never asks an open question when a menu will do, and he never leaves a cost unnamed.",
 "KEEP is a finding, not a failure. Roughly three quarters of the fields came back unchanged. Telling the writer up front that a KEEP is a legitimate answer is what stopped the pass from churning good lines into different lines, and it is what made the rewritten ones worth reading.",
 "House laws should be enforced by a checker, not trusted to the writer. The no-em-dash law and the no-simile law survived the drafting call because a second, cheaper model read every AFTER string against them. Style rules stated in a prompt are aspirations; style rules verified in a script are laws.",
 "The em dash is the sharpest conflict between his private voice and his public one. It is his native pivot in Italian and in English, in nearly every paragraph he writes. The law is right anyway: on a page the reader cannot hear his pauses, so the pivot has to become a period.",
 "Fixed UI slots are a craft constraint, not an obstacle. A plus-or-minus 25% band on every field, with headings held at or under the original length, forced the compression the copy actually needed. The constraint creed on the front page turned out to apply to the front page itself.",
 "Concreteness is the only reliable test for AI residue. Nearly every line the checker flagged failed the same way: an abstraction where a fact would fit. The repair is never a better adjective, it is the number, the street, or the year that was already available.",
 "Quotations, citations, keycodes and balloted lines need to be walled off by name. Telling the model to \"keep facts\" is not enough. Naming Welles, Hitchcock, Lumet, Chekhov, King, Morrison, Kurosawa, Glass and Stewart, and naming the heroLine ballot of 2026-09-04, is what kept them intact.",
 "Separating mechanical repair from taste repair is what let the pass close in one loop. Law breaks, length overruns and meaning drift were fixed automatically because they have right answers. Hedges, stacked adjectives and sameness were only recorded, because those are Marco's call and a second model's taste is not his.",
 "The panel has to be seated somewhere else. Claude wrote every AFTER line in this document, so Claude cannot be trusted to say whether it worked. The four chairs run on four other models on two other providers, their notes are pasted in verbatim, and nothing here softens or ranks them.",
 "Three rich-text fields in the whole site is a finding worth keeping. All but three of the 412 in-scope fields are plain strings, so the Lexical tree walk was needed exactly three times. Checking the shape of the data before designing around it saved most of the machinery this pass was expected to need.",
]

# ---------------------------------------------------------------- helpers

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:color"), "auto"); el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    trPr.append(el)

def cell_text(cell, text, size=9, bold=False, italic=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    for i, line in enumerate(str(text).split("\n")):
        r = p.add_run(("\n" if i else "") + line)
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color:
            r.font.color.rgb = color

def rows_of(doc):
    res = after[doc["key"]]
    out = []
    for f in doc["fields"]:
        if f["kind"] == "text":
            out.append((f["path"], f["before"], res[f["path"]]))
        else:
            for i, t in enumerate(f["before_texts"]):
                p = f"{f['path']}#{i}"
                out.append((p, t, res[p]))
    return out

def md_lines(document, md, base_size=10):
    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("## "):
            h = document.add_paragraph(); r = h.add_run(line[3:].strip())
            r.bold = True; r.font.size = Pt(12)
            h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(3)
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        txt = line.strip()
        if re.match(r"^\d+\.\s", txt) or txt.startswith("- "):
            p.paragraph_format.left_indent = Inches(0.25)
        for i, seg in enumerate(re.split(r"(\*\*[^*]+\*\*)", txt)):
            if not seg:
                continue
            if seg.startswith("**") and seg.endswith("**"):
                r = p.add_run(seg[2:-2]); r.bold = True
            else:
                r = p.add_run(seg)
            r.font.size = Pt(base_size)

# ---------------------------------------------------------------- build

d = Document()
sec = d.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
sec.left_margin = sec.right_margin = Inches(0.5)
sec.top_margin = sec.bottom_margin = Inches(0.55)
d.styles["Normal"].font.name = "Georgia"
d.styles["Normal"].font.size = Pt(10)

USABLE = Inches(11.69 - 1.0)

# --- title page
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run("APR 70 PICTURES"); r.bold = True; r.font.size = Pt(22)
t2 = d.add_paragraph(); r = t2.add_run("The website copy, in Marco's voice"); r.font.size = Pt(18)
t3 = d.add_paragraph(); r = t3.add_run("Writers-room compound writing pass  ·  7 September 2026")
r.font.size = Pt(11); r.italic = True

warn = d.add_paragraph()
r = warn.add_run("DRAFT ONLY. Nothing in this document has been applied. The website, the CMS and the databases were not modified in any way.")
r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

kept = sum(1 for doc in inv["docs"] for p, b, a in rows_of(doc) if a == b)
total = sum(len(rows_of(doc)) for doc in inv["docs"])

md_lines(d, f"""## What this is
The complete reader-visible English copy of staging.apr70.com, set beside a rewrite of the same copy into the voice of the founder, Marco Caruso. Every field the site shows a reader appears here exactly once, with the live text on the left and the writers-room text on the right.

## Method
Six passes, in the order used on the SilMart pass of 6 September 2026. Inventory, then promise and voice fit, then the line edit, then an AI-residue and voice check, then a reader pass, then the compound lessons. The voice profile was built by reading 35 of Marco's own sent emails (English and Italian), a filtered set of his dictated prompts, and the company creed of 20 May 2026. Every AFTER line was written by claude-opus-5 against that profile. Every AFTER line was then read back by claude-sonnet-5 against the house humanizer checklist, the voice test and the laws; mechanical defects were repaired, matters of taste were recorded and left for Marco. The critic panel that follows was run on four other models, on two other providers, because the model that wrote the draft does not get to score it.

## Scope
{inv['doc_count']} documents, {inv['field_count']} localized fields, {total} text rows, roughly {inv['word_count']:,} words. Five pages (Home, Slate, Craft, Methods, Contact), the site chrome that appears on every page, and ten property pages in slate order. The Mayors is not public and is marked private slate; it is kept last.

Of the {total} rows, {kept} came back unchanged and are shown as KEEP, and {total - kept} were rewritten. A KEEP is a decision, not a gap: the pass was instructed to return the exact original whenever a line already read as his voice, and not to churn for the sake of activity. Hero lines on the property pages were balloted separately on 4 September 2026 and were returned unchanged by instruction.

## What was excluded
Image alt text (the `media:*` keys in the snapshot, 206 fields) is not reader-visible copy. It is written for screen readers and for search, it is governed by the disclosure rules on the Methods page, and rewriting it into a founder's voice would make it worse at its job. It was left out of this pass.

## The laws that bound the rewrite
LAW 1. No similes. The strings "like a" and "like an" appear nowhere in the AFTER text.
LAW 2. No em dashes. A pause is an ellipsis or a full stop.
LAW 3. Mamet-direct. Say what a thing is.
Length. Every AFTER field stays within plus or minus 25% of its BEFORE. Headings, labels, buttons, captions, credits, scene slugs and nav items are no longer than the original, because these are fixed UI slots.

## A note on the em dash
Marco's own emails use the em dash as their main pivot, in English and in Italian, in nearly every paragraph he writes. LAW 2 overrides that habit on public copy and it is right to: in a letter the reader hears his pauses, and on a page the reader does not. On the site the pivot becomes a period, a colon or an ellipsis. En dashes inside date ranges that were already live (1950-1953, 2002-2013) were left exactly as they were.
""")

# --- spend table
h = d.add_paragraph(); r = h.add_run("Spend"); r.bold = True; r.font.size = Pt(12)
stages = {}
for e in ledger["entries"]:
    k = (e["stage"], e["model"])
    s = stages.setdefault(k, {"calls": 0, "usd": 0.0, "note": ""})
    s["calls"] += 1; s["usd"] += e["usd"]
    if "ESTIMATE" in (e.get("note") or ""):
        s["note"] = "estimate"
tbl = d.add_table(rows=1, cols=4); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, (w, txt) in enumerate([(3.2, "Stage"), (3.4, "Model"), (1.2, "Calls"), (1.6, "Cost (USD)")]):
    c = tbl.rows[0].cells[i]; c.width = Inches(w); cell_text(c, txt, 9, bold=True); shade(c, "E8E4DC")
repeat_header(tbl.rows[0])
for (stage, model), s in sorted(stages.items()):
    row = tbl.add_row().cells
    cell_text(row[0], stage); cell_text(row[1], model); cell_text(row[2], str(s["calls"]))
    cell_text(row[3], f"${s['usd']:.4f}" + ("  (est.)" if s["note"] else ""))
row = tbl.add_row().cells
cell_text(row[0], "TOTAL", 9, bold=True); cell_text(row[1], ""); cell_text(row[2], str(len(ledger['entries'])), 9, bold=True)
cell_text(row[3], f"${ledger['total_usd']:.4f}", 9, bold=True)
p = d.add_paragraph(); r = p.add_run(
    "Anthropic costs are computed from the API's own token counts at $5/M input, $25/M output, "
    "1.25x cache write and $0.50/M cache read for Opus 5, and $3/M input and $15/M output for Sonnet 5. "
    "OpenRouter costs are the figures OpenRouter returned. The xAI figure is an ESTIMATE at $3/M input "
    "and $15/M output; xAI does not return a cost with the response.")
r.font.size = Pt(8); r.italic = True

# --- voice section
d.add_page_break()
h = d.add_paragraph(); r = h.add_run("Marco's voice"); r.bold = True; r.font.size = Pt(16)
p = d.add_paragraph(); r = p.add_run(
    "Built by reading 35 of his own sent emails (2023-2026, English and Italian, roughly 12,000 words), "
    "a filtered set of his dictated prompts to agents, and the company creed of 20 May 2026. "
    "This profile is the system prompt every line in the AFTER column was written against.")
r.italic = True; r.font.size = Pt(9)
md_lines(d, voice, base_size=9.5)

# --- page sections
for doc in inv["docs"]:
    d.add_page_break()
    h = d.add_paragraph()
    r = h.add_run(doc["page"]); r.bold = True; r.font.size = Pt(16)
    sub = d.add_paragraph()
    tag = "  ·  PRIVATE SLATE, not public" if doc.get("private") else ""
    r = sub.add_run(f"{doc['route']}   ·   {doc['key']}{tag}")
    r.italic = True; r.font.size = Pt(9)
    rows = rows_of(doc)
    nk = sum(1 for _, b, a in rows if a == b)
    st = d.add_paragraph(); r = st.add_run(f"{len(rows)} rows.  {nk} KEEP.  {len(rows)-nk} rewritten.")
    r.font.size = Pt(9)
    tbl = d.add_table(rows=1, cols=3); tbl.style = "Table Grid"
    tbl.autofit = False
    widths = [Inches(2.0), Inches(4.35), Inches(4.35)]
    hdr = tbl.rows[0].cells
    for i, txt in enumerate(["Field", "BEFORE (staging.apr70.com)", "AFTER (writers room)"]):
        hdr[i].width = widths[i]; cell_text(hdr[i], txt, 9, bold=True); shade(hdr[i], "E8E4DC")
    repeat_header(tbl.rows[0])
    for path, b, a in rows:
        c = tbl.add_row().cells
        for i in range(3):
            c[i].width = widths[i]
        cell_text(c[0], path, 8)
        cell_text(c[1], b, 9)
        if a == b:
            cell_text(c[2], "KEEP (unchanged)", 9, italic=True, color=RGBColor(0x55, 0x55, 0x55))
        else:
            cell_text(c[2], a, 9)

# --- residue report summary
d.add_page_break()
h = d.add_paragraph(); r = h.add_run("The residue and voice check"); r.bold = True; r.font.size = Pt(16)
nflag = sum(len(v) for v in flags.values())
nrep = sum(len(v) for v in repairs.values())
md_lines(d, f"""Every AFTER line was read back by claude-sonnet-5 against the house humanizer checklist, the twelve-line voice test and the three laws. {nflag} lines were flagged. {nrep} were repaired by a single claude-opus-5 call per affected document, and only for mechanical defects: law breaks, length overruns and meaning drift. Flags about taste (residue vocabulary, hedges, stacked adjectives, sameness) were recorded and left alone, because those are Marco's ruling and not a second model's. The check ran once. It was not looped.

The full flag list with every suggested fix is in `docs/copy-pass/2026-09-07/residue-report.md`.""")

# --- critics
d.add_page_break()
h = d.add_paragraph(); r = h.add_run("Critic panel"); r.bold = True; r.font.size = Pt(16)
md_lines(d, """Welles, Sturges, Riskin, Capra, run in sequence against a plain-text dossier of every BEFORE and AFTER row in this document, with the voice profile summary and the constraint creed at the top.

Marco's ruling of 20 August 2026 governs this stage: **Claude writes, Claude never scores.** No Anthropic model sits on a chair here. Each note below is reproduced **verbatim**. Nothing has been summarised, softened, ranked or paraphrased.

One mechanical note. Capra's chair sits on the free rotating seat. Its first reply hit the 6,000-token ceiling and stopped mid-sentence, so the call was made again with a 14,000-token ceiling and the complete note is the one printed here. The seat, the model and the prompt were identical; only the ceiling changed.""")
for chair, name, lens in CHAIRS:
    f = OUT / "critics" / f"{chair}.md"
    d.add_page_break()
    hh = d.add_paragraph(); r = hh.add_run(name); r.bold = True; r.font.size = Pt(14)
    if not f.exists():
        p = d.add_paragraph(); r = p.add_run("Seat unavailable. No note was returned."); r.italic = True
        continue
    body = f.read_text()
    meta = ""
    m = re.match(r"<!--(.*?)-->\s*", body, re.S)
    if m:
        meta = m.group(1).strip(); body = body[m.end():]
    sp = d.add_paragraph(); r = sp.add_run(f"{lens}   ·   {meta}")
    r.italic = True; r.font.size = Pt(8)
    for para in body.split("\n"):
        if not para.strip():
            continue
        pp = d.add_paragraph(); pp.paragraph_format.space_after = Pt(3)
        r = pp.add_run(para.rstrip()); r.font.size = Pt(9.5)

# --- lessons
d.add_page_break()
h = d.add_paragraph(); r = h.add_run("Compound lessons"); r.bold = True; r.font.size = Pt(16)
p = d.add_paragraph(); r = p.add_run(
    "What this pass taught about writing in his voice. This section is written by Claude as commentary; "
    "everything above it is either the live copy, the rewrite, or a critic's note.")
r.italic = True; r.font.size = Pt(9)
for i, l in enumerate(LESSONS, 1):
    pp = d.add_paragraph(); pp.paragraph_format.left_indent = Inches(0.25)
    pp.paragraph_format.space_after = Pt(6)
    r = pp.add_run(f"{i}.  "); r.bold = True; r.font.size = Pt(10)
    r = pp.add_run(l); r.font.size = Pt(10)

DEST.parent.mkdir(parents=True, exist_ok=True)
d.save(str(DEST))
shutil.copy2(DEST, SCRATCH)
print(f"docx -> {DEST}")
print(f"docx -> {SCRATCH}")

# ---------------------------------------------------------------- markdown twin
today = "2026-09-07"
md = [f"""---
type: note
entity: APR70
project: apr70-website
status: draft
title: Website - Copy Pass 2026-09-07
date created: {today}
date updated: {today}
tags: [apr70, website, copy, writers-room, voice, draft]
---
# Website — Copy Pass 2026-09-07

**DRAFT ONLY. Nothing applied.** The website, the CMS and the databases were not modified in any way. This is the Word document's twin: `/Volumes/SharedData/10-03-website-assets/copy-pass-2026-09-07/APR70-site-copy-writers-room-2026-09-07.docx`.

## What this is

The complete reader-visible English copy of staging.apr70.com, set beside a rewrite of the same copy into Marco's own voice. {inv['doc_count']} documents, {inv['field_count']} localized fields, {total} text rows, roughly {inv['word_count']:,} words. Of those rows, **{kept} came back KEEP** (unchanged) and **{total - kept} were rewritten**.

## Method

Inventory → promise and voice fit → line edit → AI-residue and voice check → reader pass → compound lessons. Same order as the SilMart pass of 2026-09-06. The voice profile was built from 35 of Marco's own sent emails, a filtered set of his dictated agent prompts, and the company creed of 2026-05-20. `claude-opus-5` wrote every AFTER line against that profile; `claude-sonnet-5` read every AFTER line back against the humanizer checklist, the voice test and the laws; mechanical defects were repaired and matters of taste were recorded. The critic panel ran on four other models on two other providers.

## Scope and exclusions

Five pages (Home, Slate, Craft, Methods, Contact), the site chrome, and ten properties in slate order with The Mayors last and marked private slate. **Excluded: image alt text** (`media:*`, 206 fields). Alt text is not reader-visible copy; it is written for screen readers and search, and rewriting it into a founder's voice would make it worse at its job.

## The laws

- **LAW 1** — no "like a" / "like an" similes.
- **LAW 2** — **no em dashes.** A pause is an ellipsis or a full stop.
- **LAW 3** — Mamet-direct: say what a thing is.
- **Length** — every AFTER field within ±25% of its BEFORE; headings, labels and buttons no longer than the original.

**On the em dash.** Marco's own emails use the em dash as their main pivot, in English and in Italian, in nearly every paragraph. LAW 2 overrides that habit on public copy: in a letter the reader hears his pauses, on a page the reader does not. En dashes inside existing date ranges were left as they were.

## Spend
"""]
md.append("| Stage | Model | Calls | Cost (USD) |")
md.append("|---|---|---|---|")
for (stage, model), s in sorted(stages.items()):
    md.append(f"| {stage} | `{model}` | {s['calls']} | ${s['usd']:.4f}{' (est.)' if s['note'] else ''} |")
md.append(f"| **TOTAL** | | **{len(ledger['entries'])}** | **${ledger['total_usd']:.4f}** |")
md.append("\nxAI returns no cost with the response; the Riskin seat's figure is an estimate at $3/M in and $15/M out.\n")
md.append("## Marco's voice\n")
md.append(voice.strip())
md.append("\n---\n")
for doc in inv["docs"]:
    tag = "  ·  **PRIVATE SLATE, not public**" if doc.get("private") else ""
    rows = rows_of(doc)
    nk = sum(1 for _, b, a in rows if a == b)
    md.append(f"\n## {doc['page']}\n")
    md.append(f"`{doc['route']}` · `{doc['key']}`{tag} · {len(rows)} rows, {nk} KEEP, {len(rows)-nk} rewritten\n")
    md.append("| Field | BEFORE (staging.apr70.com) | AFTER (writers room) |")
    md.append("|---|---|---|")
    for path, b, a in rows:
        def esc(x):
            return str(x).replace("|", "\\|").replace("\n", "<br>")
        md.append(f"| `{path}` | {esc(b)} | {'*KEEP (unchanged)*' if a == b else esc(a)} |")
md.append("\n---\n")
md.append("## The residue and voice check\n")
md.append(f"{nflag} lines flagged, {nrep} repaired (mechanical only: law breaks, length, drift). Taste flags recorded, not applied. Full list: `docs/copy-pass/2026-09-07/residue-report.md`.\n")
md.append("## Critic panel\n")
md.append("Marco's ruling of 2026-08-20 governs this stage: **Claude writes, Claude never scores.** No Anthropic model sits on a chair. Each note below is **verbatim**. Nothing summarised, softened, ranked or paraphrased.\n\nOne mechanical note: Capra's chair sits on the free rotating seat. Its first reply hit the 6,000-token ceiling and stopped mid-sentence, so the call was made again with a 14,000-token ceiling. Same seat, same model, same prompt; only the ceiling changed.\n")
for chair, name, lens in CHAIRS:
    f = OUT / "critics" / f"{chair}.md"
    md.append(f"\n### {name}\n")
    if not f.exists():
        md.append("*Seat unavailable. No note was returned.*\n"); continue
    body = f.read_text()
    m = re.match(r"<!--(.*?)-->\s*", body, re.S)
    if m:
        md.append(f"*{lens} · {m.group(1).strip()}*\n")
        body = body[m.end():]
    md.append(body.rstrip() + "\n")
md.append("\n---\n")
md.append("## Compound lessons\n")
md.append("*Written by Claude as commentary; everything above is the live copy, the rewrite, or a critic's note.*\n")
for i, l in enumerate(LESSONS, 1):
    md.append(f"{i}. {l}\n")

VAULT.write_text("\n".join(md) + "\n", encoding="utf-8")
print(f"md   -> {VAULT}")
