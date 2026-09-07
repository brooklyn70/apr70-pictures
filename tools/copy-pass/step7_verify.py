#!/usr/bin/env python3
"""Pass 7: verify the deliverable before reporting."""
import json, pathlib, re, sys
sys.path.insert(0, str(pathlib.Path(__file__).parent))
import lib
from docx import Document

OK = True
def chk(label, cond, detail=""):
    global OK
    print(("PASS  " if cond else "FAIL  ") + label + (f"  {detail}" if detail else ""))
    if not cond:
        OK = False

inv = json.loads((lib.OUT / "inventory.json").read_text())
after = json.loads((lib.OUT / "after.json").read_text())
ledger = lib.ledger_load()

DEST = pathlib.Path("/Volumes/SharedData/10-03-website-assets/copy-pass-2026-09-07/APR70-site-copy-writers-room-2026-09-07.docx")
SCRATCH = pathlib.Path("/private/tmp/claude-501/-Users-marco-websites-apr70-website/2251e5c8-4793-420f-af9f-38fc6b2b3078/scratchpad/APR70-site-copy-writers-room-2026-09-07.docx")
VAULT = pathlib.Path("/Users/marco/vault/10 Work/11 APR70 Pictures/11.03 Company Ops/00 Website - Copy Pass 2026-09-07.md")

# --- laws on every AFTER string
em, sim, drift, longs = [], [], [], []
rows = 0
snap = json.loads((pathlib.Path(__file__).resolve().parents[2] / "docs/i18n/drafts/en.snapshot.json").read_text())
SHORT_SLOT = re.compile(r"(heading|label|Label|kicker|caption|credit|sceneSlug|term|cite|note|tagline|subtitle|metaLine|seoTitle)")
for d in inv["docs"]:
    res = after[d["key"]]
    for f in d["fields"]:
        items = [(f["path"], f["before"])] if f["kind"] == "text" else \
                [(f"{f['path']}#{i}", t) for i, t in enumerate(f["before_texts"])]
        for p, b in items:
            rows += 1
            a = res[p]
            if "—" in a:
                em.append((d["key"], p))
            if re.search(r"\blike an?\b", a):
                sim.append((d["key"], p, a))
            if a != b:
                if len(a) > len(b) * 1.25 + 2:
                    longs.append((d["key"], p, len(b), len(a)))
                # fact guard: numbers, urls, markdown links, arrows, highlight markers
                for tok in set(re.findall(r"\d{3,4}|https?://\S+|\[[^\]]+\]\([^)]+\)|→|==[^=]+==|@\S+\.\w+", b)):
                    if tok not in a:
                        drift.append((d["key"], p, tok))

chk("no em dash (U+2014) in any AFTER string", not em, str(em[:5]))
chk('no "like a"/"like an" in any AFTER string', not sim, str(sim[:3]))
chk("no AFTER string more than 25% longer than its BEFORE", not longs, str(longs[:5]))
chk("no dropped number / link / arrow / marker in rewritten fields", not drift, str(drift[:8]))

# --- coverage: every in-scope field exactly once
inscope = {(k, p) for k, v in snap.items() if not k.startswith("media:") for p in v}
covered = []
for d in inv["docs"]:
    for f in d["fields"]:
        covered.append((d["key"], f["path"]))
chk("every in-scope field present exactly once in the inventory",
    len(covered) == len(set(covered)) == len(inscope),
    f"inventory={len(covered)} unique={len(set(covered))} snapshot={len(inscope)}")
chk("after.json covers all 16 documents", len(after) == 16, f"{len(after)}")
chk("after.json key shape matches inventory for every document",
    all(set(after[d["key"]]) == {p for f in d["fields"] for p in
        ([f["path"]] if f["kind"] == "text" else [f"{f['path']}#{i}" for i in range(len(f["before_texts"]))])}
        for d in inv["docs"]))
chk("media:* excluded", not any(k.startswith("media:") for k in after))

# --- lexical reinsert sanity
rt = [(d["key"], f) for d in inv["docs"] for f in d["fields"] if f["kind"] == "richtext"]
chk("rich-text fields keep their text-node count", all(
    len(f["before_texts"]) == sum(1 for k in after[key] if k.startswith(f["path"] + "#")) for key, f in rt),
    f"{len(rt)} rich-text fields")

# --- ledger
tot = ledger["total_usd"]
chk("ledger total under the $4.50 abort line", tot < 4.50, f"${tot:.4f}")
chk("ledger total under the $5.00 hard limit", tot < 5.00, f"${tot:.4f}")
chk("no Anthropic model on a critic seat",
    not any(e["stage"].startswith("5-") and "claude" in e["model"] for e in ledger["entries"]))

# --- critics
seats = {}
for chair in ("welles", "sturges", "riskin", "capra"):
    f = lib.OUT / "critics" / f"{chair}.md"
    seats[chair] = "MISSING" if not f.exists() else (
        "SEAT UNAVAILABLE" if "SEAT UNAVAILABLE" in f.read_text() else "ok")
chk("all four critic files exist", all(v != "MISSING" for v in seats.values()), str(seats))

# --- docx
for path in (DEST, SCRATCH):
    chk(f"docx exists: {path}", path.exists(),
        f"{path.stat().st_size//1024} KB" if path.exists() else "")
if DEST.exists():
    doc = Document(str(DEST))
    tbls = doc.tables
    body_tbls = [t for t in tbls if len(t.columns) == 3]
    trows = sum(len(t.rows) - 1 for t in body_tbls)
    chk("docx has one 3-column table per document (16) plus the spend table",
        len(body_tbls) == 16 and len(tbls) == 17, f"tables={len(tbls)} 3col={len(body_tbls)}")
    chk("docx table rows equal the inventory row count", trows == rows, f"docx={trows} inventory={rows}")
    txt = "\n".join(p.text for p in doc.paragraphs)
    chk("docx carries the DRAFT ONLY warning", "DRAFT ONLY" in txt)
    chk("docx is landscape", doc.sections[0].page_width > doc.sections[0].page_height)
chk(f"markdown twin exists: {VAULT}", VAULT.exists(),
    f"{VAULT.stat().st_size//1024} KB" if VAULT.exists() else "")
if VAULT.exists():
    v = VAULT.read_text()
    chk("markdown twin has vault frontmatter", v.startswith("---\ntype: note\nentity: APR70"))
    chk("markdown twin states DRAFT ONLY", "DRAFT ONLY" in v)

print()
print(f"rows checked: {rows}   ledger: ${tot:.4f}")
print("ALL CHECKS PASS" if OK else "SOME CHECKS FAILED")
sys.exit(0 if OK else 1)
